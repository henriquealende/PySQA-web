"""Export endpoints, the reduced-signal flags and the calibrated waveform."""
import csv
import http.client
import io
import json
import os
import subprocess
import sys
import tempfile
import time
import uuid
import warnings

warnings.filterwarnings("ignore")
import numpy as np
import openpyxl
import soundfile as sf

BACKEND = os.getcwd()
fails = []


def check(name, ok, detail=""):
    print(("PASS " if ok else "FAIL ") + name + (" -> " + str(detail)[:220] if detail != "" else ""), flush=True)
    if not ok:
        fails.append(name)


def wav(seconds, sr=48000):
    t = np.arange(int(seconds * sr)) / sr
    x = 0.3 * np.sin(2 * np.pi * 1000 * t) + 0.01 * np.random.default_rng(1).standard_normal(t.size)
    b = io.BytesIO()
    sf.write(b, x.astype("float32"), sr, format="WAV", subtype="PCM_16")
    return b.getvalue()


class Server:
    def __init__(self, port, **env):
        self.port = port
        e = dict(os.environ)
        for k in ("PYSQA_MEMORY_MB", "RENDER"):
            e.pop(k, None)
        e.update({k: str(v) for k, v in env.items()})
        e["PYSQA_TMP_DIR"] = tempfile.mkdtemp(prefix="pysqa-exp-")
        e["PYTHONDONTWRITEBYTECODE"] = "1"
        self.proc = subprocess.Popen([sys.executable, "-m", "uvicorn", "main:app", "--host", "127.0.0.1", "--port", str(port)],
                                     cwd=BACKEND, env=e, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        for _ in range(60):
            try:
                if self.req("GET", "/api/health")[0] == 200:
                    return
            except Exception:
                time.sleep(1)
        raise RuntimeError("server did not start")

    def req(self, method, path, body=None, headers=None, raw=False):
        c = http.client.HTTPConnection("127.0.0.1", self.port, timeout=300)
        c.request(method, path, body=body, headers=headers or {})
        r = c.getresponse()
        data = r.read()
        c.close()
        if raw:
            return r.status, data
        try:
            return r.status, json.loads(data)
        except Exception:
            return r.status, data[:100]

    def post(self, path, obj, raw=False):
        return self.req("POST", path, json.dumps(obj), {"Content-Type": "application/json"}, raw=raw)

    def upload(self, session, data):
        b = uuid.uuid4().hex
        body = (f"--{b}\r\nContent-Disposition: form-data; name=\"session_id\"\r\n\r\n{session}\r\n"
                f"--{b}\r\nContent-Disposition: form-data; name=\"files\"; filename=\"t.wav\"\r\nContent-Type: audio/wav\r\n\r\n").encode() + data + f"\r\n--{b}--\r\n".encode()
        return self.req("POST", "/api/audio/upload", body, {"Content-Type": f"multipart/form-data; boundary={b}"})

    def stop(self):
        self.proc.terminate()
        try:
            self.proc.wait(10)
        except Exception:
            self.proc.kill()


print("== local mode ==")
s = Server(8301)
try:
    sess = "exp-" + uuid.uuid4().hex[:8]
    st, d = s.upload(sess, wav(10))
    aid = d[0]["audio_id"]

    st, w = s.req("GET", f"/api/audio/{sess}/{aid}/waveform")
    check("long waveform is flagged as reduced", st == 200 and w["reduced"] is True and len(w["t"]) <= 8000, f"reduced={w.get('reduced')} points={len(w['t'])}")
    check("waveform not calibrated yet: full scale", w["calibrated"] is False and abs(max(w["amplitude"]) - w["peak"]) < 0.01, f"max {max(w['amplitude']):.3f} peak {w['peak']:.3f}")

    st, c = s.post(f"/api/audio/{sess}/{aid}/calibrate", {"sensitivity_mv_pa": 50})
    st, w2 = s.req("GET", f"/api/audio/{sess}/{aid}/waveform")
    ratio = max(w2["amplitude"]) / max(w["amplitude"])
    check("calibrated waveform is in Pa (1000/50 = 20 Pa per full scale)", st == 200 and w2["calibrated"] is True and abs(ratio - 20.0) < 0.2, f"x{ratio:.2f}")
    check("peak stays in full scale for the clipping test", abs(w2["peak"] - w["peak"]) < 1e-9, w2["peak"])

    st, short = s.upload(sess, wav(0.1))
    st, ws = s.req("GET", f"/api/audio/{sess}/{short[0]['audio_id']}/waveform")
    check("short waveform is not reduced", ws["reduced"] is False and len(ws["t"]) == 4800, len(ws["t"]))
    st, e1 = s.req("GET", f"/api/audio/{sess}/{aid}/envelope")
    st, e2 = s.req("GET", f"/api/audio/{sess}/{short[0]['audio_id']}/envelope")
    check("envelope carries the reduced flag", e1["reduced"] is True and e2["reduced"] is False)

    payload = {"session_id": sess, "audio_ids": [aid], "plot_data": [
        {"sheet_name": "Waveform", "columns": {"Time (s)": [0.0, 0.1, 0.2], "Loudness\n(sone), ISO": [1.5, 2.5, 3.5]}},
        {"sheet_name": "Third-octave", "columns": {"Band centre (Hz)": [100, 200], "A very long header for the second column": [-20.5, None]}},
    ]}
    st, xl = s.post("/api/export/xlsx", payload, raw=True)
    wb = openpyxl.load_workbook(io.BytesIO(xl))
    ws1, ws2 = wb["Waveform"], wb["Third-octave"]
    check("xlsx has both sheets with the given names", wb.sheetnames == ["Waveform", "Third-octave"], wb.sheetnames)
    check("xlsx headers and values are right", ws1["B1"].value == "Loudness\n(sone), ISO" and ws1["B3"].value == 2.5 and ws2["B3"].value is None, [ws1["B1"].value, ws1["B3"].value])
    wa, wb2 = ws2.column_dimensions["A"].width, ws2.column_dimensions["B"].width
    check("each column is sized for its own header", wb2 > wa + 10, f"A={wa} B={wb2}")

    st, cs = s.post("/api/export/csv", {"session_id": sess, "audio_id": aid, "columns": {"Loudness\n(sone), ISO": [1.5, 2.5], "plain": [1, None]}}, raw=True)
    check("csv starts with a byte order mark", cs[:3] == b"\xef\xbb\xbf", cs[:6])
    rows = list(csv.reader(io.StringIO(cs.decode("utf-8-sig"))))
    check("csv quotes headers with commas and line breaks", rows[0] == ["Loudness\n(sone), ISO", "plain"] and rows[1] == ["1.5", "1"] and rows[2] == ["2.5", ""], rows)

    st, big = s.post("/api/export/xlsx", {"session_id": sess, "audio_ids": [], "plot_data": [{"sheet_name": "x", "columns": {"a": list(range(60000)), "b": list(range(60000))}}]}, raw=True)
    check("local mode builds a large workbook (no cap)", st == 200 and big[:2] == b"PK", st)
finally:
    s.stop()

print("== hosted mode ==")
s = Server(8302, PYSQA_MEMORY_MB=512)
try:
    ok = {"session_id": "x" * 8, "audio_ids": [], "plot_data": [{"sheet_name": "x", "columns": {"a": list(range(8000)), "b": list(range(8000)), "c": list(range(8000))}}]}
    st, d = s.post("/api/export/xlsx", ok, raw=True)
    check("a normal-sized workbook (24 000 cells) is built", st == 200 and d[:2] == b"PK", st)
    big = {"session_id": "x" * 8, "audio_ids": [], "plot_data": [{"sheet_name": "x", "columns": {"a": list(range(70000)), "b": list(range(70000))}}]}
    st, d = s.post("/api/export/xlsx", big)
    check("a huge workbook (140 000 cells) is refused and CSV is suggested", st == 413 and "CSV" in str(d) and "140,000" in str(d), d)
    st, h = s.req("GET", "/api/health")
    check("health reports the export cap", h["limits"]["max_export_cells"] == 100000.0, h["limits"].get("max_export_cells"))
finally:
    s.stop()

print("FAILS:", len(fails), fails)
sys.exit(1 if fails else 0)
