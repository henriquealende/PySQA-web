import csv
import glob
import io
import os
import sys

import openpyxl

OUT = sys.argv[1]
ok = True


def check(name, cond, detail=""):
    global ok
    ok = ok and bool(cond)
    print(("PASS " if cond else "FAIL ") + name + (" -> " + str(detail)[:200] if detail != "" else ""))


def sheet_info(path):
    wb = openpyxl.load_workbook(path)
    return {ws.title: (ws.max_row - 1, ws.max_column, [c.value for c in ws[1]][:6]) for ws in wb.worksheets}


print("files received:", sorted(os.listdir(OUT)))

# default view: waveform and spectrum
x = sheet_info(os.path.join(OUT, "export-test_waveform-and-spectrum.xlsx"))
check("waveform sheet: min/max in Pa (calibrated)", list(x["Waveform"][2]) == ["Time (s)", "Minimum (Pa)", "Maximum (Pa)"] and x["Waveform"][0] == 4000, x["Waveform"])
check("spectrum sheet: peak and average in dB SPL", x["Spectrum"][2][1].startswith("Peak (dB SPL") and x["Spectrum"][2][2].startswith("Average (dB SPL"), x["Spectrum"][2])
ws = openpyxl.load_workbook(os.path.join(OUT, "export-test_waveform-and-spectrum.xlsx"))["Waveform"]
vals_max = [ws.cell(row=r, column=3).value for r in range(2, ws.max_row + 1)]
vals_min = [ws.cell(row=r, column=2).value for r in range(2, ws.max_row + 1)]
check("clipped burst is in the data (1.0 FS x 20 Pa/FS = 20 Pa) and max >= min everywhere",
      max(vals_max) >= 19.9 and all(a >= b for a, b in zip(vals_max, vals_min)), f"max {max(vals_max):.2f} Pa")
times = [ws.cell(row=r, column=1).value for r in range(2, ws.max_row + 1)]
check("time column increases and spans the file", all(a <= b for a, b in zip(times, times[1:])) and 1.1 < times[-1] < 1.21, f"{times[0]} .. {times[-1]}")

# CSV of the chosen table
rows = list(csv.reader(io.StringIO(open(os.path.join(OUT, "export-test_spectrum.csv"), encoding="utf-8-sig").read())))
check("csv: the Spectrum table was exported (chosen in the dialog)", rows[0][0] == "Frequency (Hz)" and len(rows) - 1 == 1358, (rows[0], len(rows) - 1))
peak_col = rows[0].index([h for h in rows[0] if h.startswith("Peak")][0])
peak_row = max(rows[1:], key=lambda r: float(r[peak_col]))
check("csv: the 1 kHz tone is the strongest line", 950 < float(peak_row[0]) < 1050, peak_row[:3])

# 1/3 octave
o = sheet_info(os.path.join(OUT, "export-test_1-3-octave-bands.xlsx"))
check("octave: 31 bands with their unit", o["Third-octave"][0] == 31 and "dB SPL" in o["Third-octave"][2][1], o)
orows = list(csv.reader(open(os.path.join(OUT, "export-test_third-octave.csv"), encoding="utf-8-sig")))
lv = {float(r[0]): float(r[1]) for r in orows[1:] if r[1] != ""}
check("octave csv: the 1 kHz band is the loudest", max(lv, key=lv.get) == 1000.0, max(lv, key=lv.get))

# envelope
e = sheet_info(os.path.join(OUT, "export-test_envelope.xlsx"))
check("envelope: min/max signal and envelope plus the envelope spectrum",
      e["Envelope"][2] == ["Time (s)", "Signal minimum", "Signal maximum", "Envelope"] and "Envelope spectrum" in e, e)
erows = list(csv.reader(open(os.path.join(OUT, "export-test_envelope-spectrum.csv"), encoding="utf-8-sig")))
check("envelope csv: spectrum columns", erows[0] == ["Frequency (Hz)", "Magnitude (dB)"], erows[0])

# metric
m = sheet_info(os.path.join(OUT, "export-test_ipm.xlsx"))
check("metric: summary, by-band values (Bark HMS) and unit", set(m) == {"Summary", "By band"} and m["Summary"][2][:3] == ["Metric", "Global value", "Unit"], m)
srows = openpyxl.load_workbook(os.path.join(OUT, "export-test_ipm.xlsx"))["Summary"]
check("metric summary row: IPM, a number, dB", srows["A2"].value == "IPM" and isinstance(srows["B2"].value, float) and srows["C2"].value == "dB", [c.value for c in srows[2]])
mrows = list(csv.reader(open(os.path.join(OUT, "export-test_by-band.csv"), encoding="utf-8-sig")))
check("metric csv: 21 band rows", len(mrows) - 1 == 21, len(mrows) - 1)

# spectrogram
sp = sheet_info(os.path.join(OUT, "export-test_spectrogram-db.xlsx"))
check("spectrogram: frequency rows by time columns", sp["Spectrogram"][0] > 500 and sp["Spectrogram"][1] > 50 and sp["Spectrogram"][2][0] == "Frequency (Hz)", (sp["Spectrogram"][0], sp["Spectrogram"][1], sp["Spectrogram"][2]))
prow = list(csv.reader(open(os.path.join(OUT, "export-test_spectrogram.csv"), encoding="utf-8-sig")))
check("spectrogram csv: same shape as the sheet", len(prow) - 1 == sp["Spectrogram"][0] and len(prow[0]) == sp["Spectrogram"][1], (len(prow) - 1, len(prow[0])))

print("ALL OK" if ok else "SOME CHECKS FAILED")
sys.exit(0 if ok else 1)
