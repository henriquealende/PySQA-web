"""Export router — WAV, XLSX, CSV downloads."""
from __future__ import annotations

import csv
import io
from pathlib import Path
from typing import List

from fastapi import APIRouter, HTTPException
from fastapi.responses import FileResponse, StreamingResponse
from pydantic import BaseModel

from services.session_store import get_entry
from utils import limits
from utils.lang import msg

router = APIRouter()


@router.get("/wav/{session_id}/{audio_id}")
def export_wav(session_id: str, audio_id: str):
    entry = get_entry(session_id, audio_id)
    if not entry:
        raise HTTPException(404, msg("Áudio não encontrado", "Audio not found"))
    return FileResponse(
        path=str(entry.raw_path),
        media_type="audio/wav",
        filename=entry.name,
        headers={"Content-Disposition": f'attachment; filename="{entry.name}"'},
    )


class XLSXRequest(BaseModel):
    session_id: str
    audio_ids: List[str]
    plot_data: List[dict]   # [{sheet_name, columns: {col: [values]}, marks?: {col: ["pass"|"fail"|null]}, notes?: [text]}]


def _build_xlsx(plot_data: List[dict]) -> io.BytesIO:
    """Create a richly-formatted XLSX workbook."""
    try:
        import openpyxl
        from openpyxl.styles import (
            PatternFill, Font, Alignment, Border, Side, numbers
        )
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise HTTPException(500, msg("O pacote openpyxl não está instalado", "openpyxl not installed"))

    # ── Colour palette (matches app theme) ────────────────────────────────
    TEAL_HEX   = "FF50B89E"
    DARK_HEX   = "FF0F1923"
    PANEL_HEX  = "FF151F2B"
    ALT_HEX    = "FF1A2738"
    TEXT_WHITE = "FFE2E8F0"
    TEXT_MUTED = "FF8AB4C9"
    BORDER_HEX = "FF1E2D3D"

    header_fill = PatternFill("solid", fgColor=TEAL_HEX)
    alt_fill    = PatternFill("solid", fgColor=ALT_HEX)
    panel_fill  = PatternFill("solid", fgColor=PANEL_HEX)
    header_font = Font(bold=True, color=DARK_HEX,   size=10, name="Calibri")
    data_font   = Font(         color=TEXT_WHITE,   size=10, name="Calibri")
    muted_font  = Font(         color=TEXT_MUTED,   size=9,  name="Calibri")
    thin_border = Border(
        bottom=Side(style="thin", color=BORDER_HEX),
        right =Side(style="thin", color=BORDER_HEX),
    )
    center = Alignment(horizontal="center", vertical="center")
    right  = Alignment(horizontal="right",  vertical="center")
    left   = Alignment(horizontal="left",   vertical="center")

    # Cells marked "pass" or "fail" (batch limits): a green or red fill with bold text
    mark_fill = {"pass": PatternFill("solid", fgColor="FF1F5A3E"), "fail": PatternFill("solid", fgColor="FF6B2420")}
    mark_font = Font(bold=True, color=TEXT_WHITE, size=10, name="Calibri")

    NUM_FMT  = '#,##0.000'    # 3 decimal places with thousands sep
    INT_FMT  = '#,##0'

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    for sheet_info in plot_data:
        sheet_name = str(sheet_info.get("sheet_name", "Sheet"))[:31]
        columns: dict = sheet_info.get("columns", {})
        if not columns:
            continue
        marks: dict = sheet_info.get("marks") or {}
        notes = sheet_info.get("notes") or []

        ws = wb.create_sheet(title=sheet_name)
        ws.sheet_view.showGridLines = False
        ws.sheet_properties.tabColor = "50B89E"

        col_names = list(columns.keys())
        n_cols    = len(col_names)
        n_rows    = max((len(v) for v in columns.values()), default=0)

        # ── Header row ─────────────────────────────────────────────────
        for ci, col_name in enumerate(col_names, start=1):
            # Use col_name as-is if it contains \n (pre-formatted with unit);
            # otherwise apply title-case formatting
            display = col_name if ("\n" in col_name or " " in col_name) else col_name.replace("_", " ").title()
            cell = ws.cell(row=1, column=ci, value=display)
            cell.fill      = header_fill
            cell.font      = header_font
            has_wrap       = "\n" in display
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=has_wrap)
            cell.border    = thin_border

        # ── Data rows ──────────────────────────────────────────────────
        for ri in range(n_rows):
            row_fill = panel_fill if ri % 2 == 0 else alt_fill
            for ci, col_name in enumerate(col_names, start=1):
                vals = columns[col_name]
                val  = vals[ri] if ri < len(vals) else None
                cell = ws.cell(row=ri + 2, column=ci, value=val)
                cell.fill   = row_fill
                cell.border = thin_border

                if isinstance(val, float):
                    cell.font           = data_font
                    cell.number_format  = NUM_FMT
                    cell.alignment      = right
                elif isinstance(val, int):
                    cell.font           = data_font
                    cell.number_format  = INT_FMT
                    cell.alignment      = right
                else:
                    cell.font      = data_font
                    cell.alignment = left

                col_marks = marks.get(col_name)
                mark = col_marks[ri] if isinstance(col_marks, list) and ri < len(col_marks) else None
                if mark in mark_fill:
                    cell.fill = mark_fill[mark]
                    cell.font = mark_font

        # ── Column widths (auto-fit heuristic) ─────────────────────────
        for ci, col_name in enumerate(col_names, start=1):
            shown = col_name if ("\n" in col_name or " " in col_name) else col_name.replace("_", " ").title()
            first_line = shown.split("\n")[0]
            header_len = len(first_line) + 2
            max_content = 0
            for ri in range(min(n_rows, 500)):   # sample first 500 rows
                v = columns[col_name][ri] if ri < len(columns[col_name]) else None
                if v is not None:
                    max_content = max(max_content, len(str(v)))
            width = min(max(header_len, max_content + 2), 30)
            ws.column_dimensions[get_column_letter(ci)].width = width

        # ── Header row height (taller for multi-line headers) ──
        ws.row_dimensions[1].height = 32 if any("\n" in c for c in col_names) else 20

        # ── Notes under the table (the limits used, the count of passes) ────
        for ni, text in enumerate(notes):
            c = ws.cell(row=n_rows + 3 + ni, column=1, value=str(text))
            c.font = muted_font
            c.alignment = left

        # ── Freeze header row ──────────────────────────────────────────
        ws.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


@router.post("/xlsx")
def export_xlsx(body: XLSXRequest):
    cells = sum(len(v) for sheet in body.plot_data for v in (sheet.get("columns") or {}).values())
    limits.check_export_size(cells)
    buf = _build_xlsx(body.plot_data)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="pysqa_export.xlsx"'},
    )


class CSVRequest(BaseModel):
    session_id: str
    audio_id: str
    columns: dict


@router.post("/csv")
def export_csv(body: CSVRequest):
    out = io.StringIO()
    writer = csv.writer(out)               # quotes headers that contain commas or line breaks
    writer.writerow(body.columns.keys())
    rows = max((len(v) for v in body.columns.values()), default=0)
    for i in range(rows):
        writer.writerow(["" if i >= len(vals) or vals[i] is None else vals[i] for vals in body.columns.values()])
    return StreamingResponse(
        io.BytesIO(out.getvalue().encode("utf-8-sig")),      # BOM so Excel reads accents correctly
        media_type="text/csv",
        headers={"Content-Disposition": 'attachment; filename="pysqa_export.csv"'},
    )
